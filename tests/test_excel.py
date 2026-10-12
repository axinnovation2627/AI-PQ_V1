from openpyxl import load_workbook

from pqweb.excel_export import export
from pqweb.models import CareerRow, CertInfo, CheckReport
from pqweb.rules import judge_all


def test_export_layout(tmp_path, crit):
    rows = [
        CareerRow("2020.03.03", "2020.10.28", 30, "2020년 수도권본부 비탈면옹벽 정밀안전점검 및 성능평가용역",
                  발주자="한국도로공사수도권본부", 공사종류="고속국도", 담당업무="정밀안전점검및성능평가",
                  공사개요="대상시설물(2종시설물): 비탈면(27개소)", 쪽=62),
        CareerRow("2021.03.04", "2021.05.04", 17, "국도35호선 삼척 갈전지구 외 5개소 절토사면 정밀안전점검용역",
                  발주자="정선국토관리사무소", 공사종류="국도", 쪽=65),
        CareerRow("2019.08.12", "2019.09.02", 3, "국도46호선 당림교 정밀안전진단 및 성능평가용역",
                  공사종류="장대교량(100m이상)", 쪽=58),
    ]
    js = judge_all(rows, crit)
    rep = CheckReport(실패쪽=[3])
    out = export({"성명": "기술인A", "소속": "(주)가나엔지니어링", "참여형태": "분야책임", "모델": "fake"},
                 CertInfo(등급="토목 특급"), rows, js, rep, crit, tmp_path / "t.xlsx")
    ws = load_workbook(out).active
    # A~O: AI 판정 (C = 공사(용역)개요)
    assert ws["B8"].value == "사 업 명" and ws["C8"].value == "공사(용역)개요"
    assert ws["C10"].value == "대상시설물(2종시설물): 비탈면(27개소)"
    assert ws["F9"].value == "인정일" and ws["F10"].value == 30
    assert ws["L8"].value.startswith("AI 인정여부") and ws["O8"].value == "AI 판단 근거"
    assert [ws[f"L{r}"].value for r in (10, 11, 12)] == ["유사", "유사", "불인정"]
    assert ws["M10"].value == 1.0 and ws["M11"].value == 0.5 and ws["M12"].value is None
    assert ws["N10"].value == '=IF(M10="","",F10*M10)'
    assert ws["F13"].value == '=SUMIF(L10:L12,"유사",F10:F12)'
    assert ws["N13"].value == '=SUMIF(L10:L12,"유사",N10:N12)'
    assert ws["C5"].value == "=F13" and ws["E5"].value == "=N13"
    # Q~T: PQ 평가위원 평가 (P 공백)
    assert ws["P8"].value is None
    assert ws["Q8"].value == "인정여부" and ws["T8"].value == "평가위원 판단사유"
    assert ws["S10"].value == '=IF(R10="","",F10*R10)'
    assert ws["S13"].value == '=SUMIF(Q10:Q12,"유사",S10:S12)'
    assert "p03" in ws["L3"].value
    assert ws["I4"].value == "분야책임 ◀"
    texts = [ws.cell(r, 1).value or "" for r in range(17, ws.max_row + 1)]
    assert any("2종 증빙" in t for t in texts)


def test_current_workplace():
    from pqweb.excel_export import current_workplace
    hist = ["1998.03.02 ~ 2013.12.31 (주)다라건설", "2014.01.06 ~ 근무중 (주)가나엔지니어링",
            "1985.03.04 ~ 1989.12.30 (주)마바 現:(주)마바엔지니어링"]
    assert current_workplace(hist) == "2014.01.06 ~ 근무중 (주)가나엔지니어링"
    # '근무중' 표시가 없으면 시작일이 가장 늦은 곳
    assert current_workplace(["2001.01.01 ~ 2003.01.01 A", "2010.05.01 ~ 2012.01.01 B"]) == "2010.05.01 ~ 2012.01.01 B"
    assert current_workplace([]) == "-"


def test_export_shows_only_current_workplace(tmp_path, crit):
    info = CertInfo(근무처=["1998.03.02 ~ 2013.12.31 (주)다라건설", "2014.01.06 ~ 근무중 (주)가나엔지니어링"])
    ws = load_workbook(export({}, info, [], [], CheckReport(), crit, tmp_path / "w.xlsx")).active
    assert ws["B3"].value == "현 근무처" and ws["C3"].value == "2014.01.06 ~ 근무중 (주)가나엔지니어링"


def test_export_falls_back_when_locked(tmp_path, crit, monkeypatch):
    from openpyxl.workbook import workbook as wbmod
    real = wbmod.Workbook.save
    calls = []

    def save(self, path):
        calls.append(str(path))
        if len(calls) == 1:
            raise PermissionError
        return real(self, path)

    monkeypatch.setattr(wbmod.Workbook, "save", save)
    out = export({}, CertInfo(), [], [], CheckReport(), crit, tmp_path / "t.xlsx")
    assert out.name == "t_2.xlsx" and out.exists()
