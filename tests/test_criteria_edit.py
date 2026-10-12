import copy
import shutil

import pytest
from openpyxl import load_workbook

from pqweb import criteria as C
from pqweb.excel_export import export, role_formula
from pqweb.models import CareerRow, CertInfo, CheckReport
from pqweb.rules import judge_row


@pytest.fixture
def cfile(tmp_path):
    p = tmp_path / "2026_비탈면옹벽.yaml"
    shutil.copy(C.DEFAULT, p)
    return p


@pytest.mark.parametrize("change, msg", [
    (lambda c: c.update(기산일="1995-01-05"), "기산일"),
    (lambda c: c["대상어휘"].append(""), "대상어휘"),
    (lambda c: c.update(고속노선정규식="(경부"), "정규식"),
    (lambda c: c["인정율"]["유사"].update(점검=1.5), "인정율"),
    (lambda c: c["배점"]["분야참여"].insert(0, [2, 9.0]), "줄어들어야"),
    (lambda c: c["배점"]["사업책임"].pop(), "0년"),
    (lambda c: c.update(일수_연환산=0), "환산"),
])
def test_validate_rejects(crit, change, msg):
    c = copy.deepcopy(crit)
    change(c)
    with pytest.raises(ValueError, match=msg):
        C.validate(c)


def test_save_backup_history_restore(cfile, crit):
    c = copy.deepcopy(crit)
    c["대상어휘"].append("법면")
    b1 = C.save_criteria(c, cfile, note="법면 추가")
    assert b1 and b1.exists()
    assert "법면" in C.load_criteria(cfile)["대상어휘"]
    h = C.history(cfile)
    assert len(h) == 1 and h[0]["메모"] == "법면_추가"
    back = C.restore(h[0]["파일"], cfile)
    assert "법면" not in back["대상어휘"] and "법면" not in C.load_criteria(cfile)["대상어휘"]
    assert len(C.history(cfile)) == 2          # 되돌리기 직전 상태도 백업
    with pytest.raises(ValueError):
        C.restore("../x.yaml", cfile)


def test_edited_words_change_judgement(crit, row):
    r = row(사업명="국도38호선 법면 정밀점검용역", 발주자="정선국토관리사무소", 공사종류="")
    assert judge_row(r, crit).인정여부 == "불인정"
    c = copy.deepcopy(crit)
    c["대상어휘"].append("법면")
    assert judge_row(r, c).인정여부 == "유사"


def test_role_formula_follows_bands():
    assert role_formula([[3, 1.5], [2, 1.35], [0, 0.9]], "G4") == "IF(G4>=3,1.5,IF(G4>=2,1.35,0.9))"
    assert role_formula([[0, 2]], "G4") == "2"


def test_excel_uses_edited_scores(tmp_path, crit):
    c = copy.deepcopy(crit)
    c["배점"]["분야참여"] = [[2, 3.0], [0, 1.0]]
    c["인정율"]["유사"]["점검"] = 0.4
    rows = [CareerRow("2021.03.04", "2021.05.04", 17, "국도35호선 절토사면 정밀안전점검용역", 발주자="정선국토관리사무소")]
    from pqweb.rules import judge_all
    js = judge_all(rows, c)
    assert js[0].인정율 == 0.4
    ws = load_workbook(export({}, CertInfo(), rows, js, CheckReport(), c, tmp_path / "x.xlsx")).active
    assert ws["J5"].value == "=IF(G4>=2,3,1)"
    assert "점검 40%" in ws["K5"].value
