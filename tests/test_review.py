import json

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from pqweb.server import app

PAGES = [
    {"쪽종류": "기술경력", "경력": [
        {"시작일": "2021.03.04", "종료일": "2021.05.04", "인정일": 20, "공사종류": "국도",
         "사업명": "국도35호선 삼척 갈전지구 외 5개소 절토사면 정밀안전점검용역", "발주자": "정선국토관리사무소"},
        {"시작일": "2022.03.29", "종료일": "2022.06.26", "인정일": 6, "공사종류": "장대교량(100m이상)",
         "사업명": "국도 38호선 연당2교(하) 등 9개 시설물 정밀안전점검용역"}]},
    {"쪽종류": "기술경력", "경력": [
        {"시작일": "2023.03.15", "종료일": "2023.12.15", "인정일": 25, "공사종류": "고속국도",
         "사업명": "2023년 충북본부 깎기비탈면 옹벽 안전점검 및 성능평가용역", "발주자": "한국도로공사충북본부",
         "담당업무": "안전점검", "공사개요": "깎기비탈면 58개소"}]},
]


@pytest.fixture
def done(tmp_path, monkeypatch):
    # 가짜 판독은 호출 순서대로 쪽 내용을 돌려주므로 병렬 판독을 끈다(.env의 GEMINI_WORKERS 무시).
    monkeypatch.setenv("GEMINI_WORKERS", "1")
    app.state.base = tmp_path
    app.state.run_async = False
    it = iter(PAGES)
    app.state.make_caller = lambda: (lambda img, prompt: json.dumps(next(it), ensure_ascii=False))
    doc = pymupdf.open()
    for _ in PAGES:
        doc.new_page()
    c = TestClient(app)
    eid = c.post("/api/evaluations", files={"file": ("a.pdf", doc.tobytes(), "application/pdf")},
                 data={"name": "홍o동", "role": "분야참여"}).json()["id"]
    return c, eid


def test_rows_endpoint_gives_stable_keys_and_ai(done):
    c, eid = done
    d = c.get(f"/api/evaluations/{eid}/rows").json()
    assert [r["key"] for r in d["rows"]] == ["p01-1", "p01-2", "p02-1"]
    r0 = d["rows"][0]
    assert r0["ai"]["인정여부"] == "유사" and r0["ai"]["인정율"] == 0.5
    assert r0["사업명"].startswith("국도35호선") and d["rows"][2]["공사개요"] == "깎기비탈면 58개소"
    assert d["review"] == {"평가위원": "", "rows": {}}
    assert d["배점"]["분야참여"][0] == [3, 1.5]
    assert d["일반"]["성명"] == "홍o동" and d["일반"]["참여형태"] == "분야참여"
    assert any(m.startswith("[판정 방법]") for m in d["메모"])
    assert "직무분야" in d["rows"][0]


def test_save_and_reload_review(done):
    c, eid = done
    body = {"평가위원": "김o수", "rows": {
        "p01-1": {"인정여부": "유사", "인정율": 0.5, "사유": ""},
        "p02-1": {"인정여부": "유사", "인정율": 0.5, "사유": "담당업무 안전점검 — 성능평가 미명시"}}}
    assert c.put(f"/api/evaluations/{eid}/review", json=body).status_code == 200
    got = c.get(f"/api/evaluations/{eid}/rows").json()["review"]
    assert got["평가위원"] == "김o수" and got["rows"]["p02-1"]["인정율"] == 0.5


@pytest.mark.parametrize("bad", [
    {"rows": {"p01-1": {"인정여부": "모름"}}},
    {"rows": {"p01-1": {"인정율": 1.5}}},
    {"rows": {"p99-9": {"인정여부": "유사"}}},
])
def test_review_validation(done, bad):
    c, eid = done
    assert c.put(f"/api/evaluations/{eid}/review", json={"평가위원": "", **bad}).status_code == 400


def test_excel_with_review(done):
    c, eid = done
    c.put(f"/api/evaluations/{eid}/review", json={"평가위원": "김o수", "rows": {
        "p01-2": {"인정여부": "불인정", "인정율": None, "사유": "교량"},
        "p02-1": {"인정여부": "유사", "인정율": 0.5, "사유": "담당업무 점검"}}})
    plain = c.get(f"/api/evaluations/{eid}/excel")
    rev = c.get(f"/api/evaluations/{eid}/excel?review=1")
    assert plain.status_code == rev.status_code == 200
    assert "filename*=utf-8''" in rev.headers["content-disposition"]
    p = c.app.state.base / eid
    ws = load_workbook(sorted(p.glob("*평가위원*.xlsx"))[0]).active
    assert ws["Q10"].value is None                      # 평가 안 한 행은 빈칸
    assert ws["Q11"].value == "불인정" and ws["T11"].value == "교량"
    assert ws["Q12"].value == "유사" and ws["R12"].value == 0.5
    assert ws["Q12"].fill.fgColor.rgb.endswith("F8CBAD")   # AI(100%)와 인정율만 달라도 빨간 표시
    assert ws["Q11"].fill.fgColor.rgb in ("00000000", None)  # AI와 같음(불인정) → 표시 없음
    assert "김o수" in ws["Q4"].value
    ws_plain = load_workbook(p / json.loads((p / "status.json").read_text(encoding="utf-8"))["결과"]["엑셀"]).active
    assert ws_plain["Q11"].value is None


def test_page_image(done):
    c, eid = done
    r = c.get(f"/api/evaluations/{eid}/pages/2")
    assert r.status_code == 200 and r.headers["content-type"] == "image/png"
    assert c.get(f"/api/evaluations/{eid}/pages/99").status_code == 404
