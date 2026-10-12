"""평가시트(엑셀) 생성.

A~O: AI 판정(C: 공사(용역)개요), P: 공백, Q~T: PQ 평가위원 평가(사람이 AI 판단을 검증하는 칸).
소계·환산·점수는 엑셀 수식이라 평가위원이 값을 바꾸면 바로 다시 계산된다.
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .models import CareerRow, CertInfo, CheckReport, Judgement

FONT = "맑은 고딕"
_thin = Side(style="thin", color="808080")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
FILL = {k: PatternFill("solid", fgColor=v) for k, v in {
    "gray": "E7E6E6", "blue": "DDEBF7", "orange": "FCE4D6", "yellow": "FFF2CC",
    "green": "E2EFDA", "ai": "FFFF00", "red": "F8CBAD"}.items()}
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

HEAD = ["용역유형\n(사업명 판정)", "사 업 명", "공사(용역)개요", "시작일", "종료일", "인정일", "발주청\n(관리주체)",
        "공사종류", "직무\n분야", "전문분야", "담당업무", "AI 인정여부\n판단", "인정율", "환산일수", "AI 판단 근거"]
AI_COLS = "ABCDEFGHIJKLMNO"        # A~O: AI 판정
REVIEW_COLS = "QRST"               # P 공백, Q~T: PQ 평가위원 평가
WIDTH = dict(zip("ABCDEFGHIJKLMNOPQRST",
                 (12, 40, 40, 10, 10, 7, 20, 11, 7, 13, 18, 10, 8, 9, 56, 3, 11, 9, 10, 40)))
def _num(x) -> str:
    return f"{x:g}"


def role_formula(bands: list, y: str) -> str:
    """배점표 [[기준년수, 점수], ...](내림차순, 마지막 0년)를 엑셀 IF 수식으로."""
    *upper, (_, last) = bands
    out = _num(last)
    for th, pts in reversed(upper):
        out = f"IF({y}>={_num(th)},{_num(pts)},{out})"
    return out


def rate_text(crit: dict) -> str:
    r = crit["인정율"]
    pc = lambda v: f"{round(v * 100):g}%"
    return (f"유사: 진단·성능평가 {pc(r['유사']['진단성능'])}, 점검 {pc(r['유사']['점검'])} / "
            f"관련: 진단·성능평가 {pc(r['관련']['진단성능'])}, 점검 {pc(r['관련']['점검'])}")


class _Sheet:
    def __init__(self, ws):
        self.ws = ws

    def cell(self, ref, value=None, fill=None, bold=False, align=CENTER, fmt=None, size=10, border=True):
        c = self.ws[ref]
        if value is not None:
            c.value = value
        c.font = Font(name=FONT, size=size, bold=bold)
        c.alignment = align
        if border:
            c.border = BORDER
        if fill:
            c.fill = FILL[fill]
        if fmt:
            c.number_format = fmt
        return c

    def box(self, rng, value=None, **kw):
        first, last = (rng.split(":") + [rng])[:2]
        for row in self.ws[first:last]:
            for c in row:
                self.cell(c.coordinate, **{k: v for k, v in kw.items() if k != "fmt"})
        if ":" in rng:
            self.ws.merge_cells(rng)
        return self.cell(first, value, **kw)


def current_workplace(history: list[str]) -> str:
    """근무처 이력 중 현 근무처: '근무중' 표시가 있는 곳, 없으면 시작일이 가장 늦은 곳."""
    if not history:
        return "-"
    for h in history:
        if "근무중" in h or "현재" in h:
            return h
    return max(history, key=lambda h: (re.match(r"\s*(\d{4}\D\d{1,2}\D\d{1,2})", h) or [""])[0])


def _date(s: str):
    try:
        return datetime.strptime(s, "%Y.%m.%d")
    except (TypeError, ValueError):
        return s


def build_notes(meta: dict, judgements: list[Judgement], check: CheckReport, crit: dict) -> list[str]:
    """판정 방법·기준·확인필요·판독 검증 메모. 엑셀 하단과 평가위원 화면이 같은 내용을 쓴다.
    '['로 시작하는 줄은 소제목이다."""
    g = crit["근거"]
    flags = Counter(x.split(":")[0].split(" —")[0] for j in judgements for x in j.확인필요)
    has_inspection = any(j.분야 == "정밀점검" and j.인정여부 != "불인정" for j in judgements)
    notes = [
        f"[판정 방법] 평가기준: {crit['이름']} — AI 인정여부·인정율 모두 사업명 기준, 담당업무는 참고",
        "· 정의 — ① 해당분야: " + crit["정의"]["해당분야"] + "  ② 용역유형: " + crit["정의"]["용역유형"]["진단성능"] + " / " + crit["정의"]["용역유형"]["점검"] + "  ③ 도로유형: 당해분야(" + crit["정의"]["도로유형"]["당해분야"] + ") → 유사용역 / 그 외 분야(" + crit["정의"]["도로유형"]["그외분야"] + ") → 관련용역",
        "· 판정 순서: R0 기간(기산일·준공) → R1 경력 구분(감리 제외) → R2 용역 종류 → R3 대상 시설(절토사면·옹벽) "
        "→ R4 위치(도로법 제10조 도로면 유사, 그 외 관련) → R5 인정율",
        f"· {g['용역']}", f"· {g['유사']}", f"· {g['관련']}", f"· {g['인정율']}", f"· 현재 적용 인정율 — {rate_text(crit)}",
        f"[확인필요] 노란 칸 {sum(1 for j in judgements if j.확인필요)}건 — 평가위원 확인 필요",
    ]
    notes += [f"· {k}: {v}건" for k, v in flags.most_common()]
    if has_inspection:
        notes.append(f"· 2종 증빙: {g['2종증빙']} — 이 경력증명서만으로는 확인 불가")
    notes.append("[판독 검증] 경력증명서 '공사종류별 인정일수 현황'과 판독 합계 대조 "
                 "(0 ≤ 증명서−판독 ≤ 행 수이면 정상: 중복기간 소수점 절삭)")
    if check.대조가능:
        notes += [f"· {c['공사종류']}: 판독 {c['판독']:,} / 증명서 {c['증명서']:,}일 (행 {c['행수']}) — {c['상태']}"
                  for c in check.대조]
    else:
        notes.append("· 합계표를 찾지 못해 대조하지 못함")
    notes += [f"· 형식 오류: {x}" for x in check.형식오류]
    if check.실패쪽:
        notes.append("· 판독 실패 쪽: " + ", ".join(f"p{p:02d}" for p in check.실패쪽))
    notes.append(f"· 판독 모델: {meta.get('모델', '-')} / 생성: {meta.get('생성', '-')}")
    return notes


def reading_status(check: CheckReport) -> str:
    """판독 상태 한 줄 요약. 문제가 없으면 빈 문자열."""
    warn = []
    if check.실패쪽:
        warn.append("판독 실패 쪽 " + ", ".join(f"p{p:02d}" for p in check.실패쪽) + " — 수동 확인 필요")
    if not check.대조가능:
        warn.append("증명서 합계표 없음 — 판독 대조 불가")
    elif any(c["상태"] != "정상" for c in check.대조):
        warn.append("합계 대조 불일치 — 하단 판독 검증 참조")
    if check.형식오류:
        warn.append(f"형식 오류 {len(check.형식오류)}건")
    return " / ".join(warn)


def export(meta: dict, info: CertInfo, rows: list[CareerRow], judgements: list[Judgement],
           check: CheckReport, crit: dict, out_path: Path,
           review: list[dict | None] | None = None, reviewer: str = "") -> Path:
    """review: 행 순서에 맞춘 평가위원 입력({인정여부, 인정율, 사유} 또는 None). 주면 Q~T열을 채운다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "경력평가"
    s = _Sheet(ws)
    for col, w in WIDTH.items():
        ws.column_dimensions[col].width = w

    r0, last = 10, 10 + len(rows) - 1
    sub = last + 1
    role = meta.get("참여형태") or ""

    # ── 상단: 일반사항 ──
    ws["A1"] = "나. 참여기술인 경력 — AI 평가시트"
    ws["A1"].font = Font(name=FONT, size=13, bold=True)
    s.box("A2:A3", "일반\n사항", fill="gray", bold=True)
    s.cell("B2", "참여형태", fill="gray")
    s.box("C2:D2", role or "(미지정 — 3개 참여형태 병기)")
    s.cell("E2", "성 명", fill="gray")
    s.box("F2:G2", meta.get("성명") or info.성명)
    s.cell("H2", "소 속", fill="gray")
    s.box("I2:J2", meta.get("소속", ""))
    s.cell("K2", "등 급", fill="gray")
    s.box("L2:O2", info.등급)
    s.cell("B3", "현 근무처", fill="gray")
    s.box("C3:J3", current_workplace(info.근무처), align=LEFT, size=9)
    s.cell("K3", "판독", fill="gray")
    warn = reading_status(check)
    s.box("L3:O3", warn or "정상(합계 대조 일치)", fill="red" if warn else "green",
          align=LEFT, size=9)

    # ── 경력 요약 (수식: F=인정일, N=환산일수 소계) ──
    s.box("A4:A6", "경력\n요약", fill="gray", bold=True)
    for ref, t in (("B4", "합    계"), ("B5", "유사용역 경력"), ("B6", "관련과업 경력")):
        s.cell(ref, t, fill="gray")
    s.cell("C4", "인정일", fill="gray")
    s.cell("D4", "인정율", fill="gray")
    s.box("E4:F4", "환산일수", fill="gray")
    s.cell("G4", "환산년수", fill="gray")
    for col, rl in zip(("H", "I", "J"), ("사업책임", "분야책임", "분야참여")):
        s.cell(f"{col}4", rl + (" ◀" if rl == role else ""), fill="gray", bold=rl == role)
    s.box("K4:O4", "비고", fill="gray")
    for r, src in ((5, sub), (6, sub + 1)):
        s.cell(f"C{r}", f"=F{src}", fmt='#,##0"일"')
        s.cell(f"D{r}", f"=IF(C{r}=0,0,E{r}/C{r})", fmt="0.0%")
        s.box(f"E{r}:F{r}", f"=N{src}", fmt='#,##0.0"일"')
        s.cell(f"G{r}", f"=E{r}/{crit['일수_연환산']}", fmt='0.00"년"')
    s.cell("C4", "=C5+C6", fmt='#,##0"일"', bold=True)
    s.cell("D4", "=IF(C4=0,0,E4/C4)", fmt="0.0%")
    s.box("E4:F4", "=E5+E6", fmt='#,##0.0"일"', bold=True)
    s.cell("G4", f"=E4/{crit['일수_연환산']}", fmt='0.00"년"', bold=True)
    for col, rl in zip(("H", "I", "J"), ("사업책임", "분야책임", "분야참여")):
        s.box(f"{col}5:{col}6", "=" + role_formula(crit["배점"][rl], "G4"), fill="green", bold=True,
              fmt='0.00"점"', size=12)
    s.box("K5:O6", rate_text(crit) + "\n"
                   f"일→년 환산 {crit['일수_연환산']}일(원문 기준 없음) / 이적 보정계수 미적용(현 근무처 참고)",
          align=LEFT, size=9)
    for r in range(2, 7):
        ws.row_dimensions[r].height = 26

    # ── 경력 표 머리글 ──
    s.box("A7:O7", "경 력 사 항  (AI 판정)", fill="gray", bold=True, size=11)
    for col, t in zip(AI_COLS, HEAD):
        if col in "DEF":
            continue
        s.box(f"{col}8:{col}9", t, fill="ai" if t.startswith("AI") else "gray", bold=t.startswith("AI"))
    s.box("D8:F8", "참여기간", fill="gray")
    for col, t in zip("DEF", ("시작일", "종료일", "인정일")):
        s.cell(f"{col}9", t, fill="gray")
    ws.row_dimensions[9].height = 28

    # ── 경력 행 ──
    for i, (row, j) in enumerate(zip(rows, judgements)):
        n = r0 + i
        base = {"유사": "blue", "관련": "orange"}.get(j.인정여부)
        flag = "yellow" if j.확인필요 else base
        vals = [j.용역유형 or "-", row.사업명, row.공사개요, _date(row.시작일), _date(row.종료일), row.인정일, row.발주자,
                row.공사종류, row.직무분야, row.전문분야, row.담당업무, j.인정여부]
        for col, v in zip("ABCDEFGHIJKL", vals):
            s.cell(f"{col}{n}", v, fill=base, align=LEFT if col in "BC" else CENTER,
                   size=9 if col == "C" else 10,
                   fmt="yy/mm/dd" if col in "DE" else ("#,##0" if col == "F" else None),
                   bold=col == "L" and j.인정여부 != "불인정")
        s.cell(f"M{n}", j.인정율, fill=flag, fmt="0%")
        s.cell(f"N{n}", f'=IF(M{n}="","",F{n}*M{n})', fill=base, fmt="#,##0.0")
        why = j.근거 + "".join(f"\n[확인필요] {x}" for x in j.확인필요)
        s.cell(f"O{n}", why, fill=flag, align=LEFT, size=9)
        ws.row_dimensions[n].height = 33 + 15 * len(j.확인필요)
    rng = lambda c: f"{c}{r0}:{c}{last}"

    # ── 소계 ──
    for k, (label, crit_word) in enumerate((("유사용역", "유사"), ("관련용역", "관련"))):
        n = sub + k
        s.cell(f"A{n}", label, fill="gray", bold=True)
        s.box(f"B{n}:E{n}", "소   계", fill="gray")
        s.cell(f"F{n}", f'=SUMIF({rng("L")},"{crit_word}",{rng("F")})', fill="green", bold=True, fmt="#,##0")
        s.box(f"G{n}:M{n}", None, fill="green")
        s.cell(f"N{n}", f'=SUMIF({rng("L")},"{crit_word}",{rng("N")})', fill="green", bold=True,
               fmt="#,##0.0")
        s.cell(f"O{n}", None, fill="green")
    n = sub + 2
    s.cell(f"A{n}", None, fill="gray")
    s.box(f"B{n}:E{n}", "총   계", fill="gray", bold=True)
    s.cell(f"F{n}", f"=F{sub}+F{sub + 1}", fill="green", bold=True, fmt="#,##0")
    s.box(f"G{n}:M{n}", None, fill="green")
    s.cell(f"N{n}", f"=N{sub}+N{sub + 1}", fill="green", bold=True, fmt="#,##0.0")
    s.cell(f"O{n}", None, fill="green")

    # ── PQ 평가위원 평가 (Q 인정여부, R 인정율, S 환산일수, T 판단사유) ──
    s.cell("Q2", "사업책임", fill="gray")
    s.cell("R2", "분야책임", fill="gray")
    s.cell("S2", "분야참여", fill="gray")
    s.box("T2:T3", "평가위원 점수 (환산년수 기준 자동 계산)", align=LEFT, size=9)
    for col, rl in zip("QRS", ("사업책임", "분야책임", "분야참여")):
        s.cell(f"{col}3", '=IF(R7="","",' + role_formula(crit["배점"][rl], "R7") + ")", bold=True, fmt='0.00"점"')
    s.box("Q4:T5", f"PQ 평가위원 평가\n(평가위원: {reviewer or '             '})", bold=True, size=11)
    s.cell("Q6", "환산일수", fill="gray")
    s.box("R6:S6", f"=S{sub + 2}", fmt='#,##0.0"일"')
    s.cell("Q7", "환산년수", fill="gray")
    s.box("R7:S7", f'=IF(COUNT({rng("S")})=0,"",S{sub + 2}/{crit["일수_연환산"]})', fmt='0.00"년"')
    s.box("T6:T7", "※ 인정여부·인정율 입력 시 자동 계산\n※ AI 판단과 다르면 판단사유 기재", align=LEFT, size=9)
    for col, t in zip(REVIEW_COLS, ("인정여부", "인정율", "환산일수", "평가위원 판단사유")):
        s.box(f"{col}8:{col}9", t, fill="gray", bold=col in "QT")
    for i, n in enumerate(range(r0, last + 1)):
        v = (review[i] if review and i < len(review) else None) or {}
        differs = bool(v.get("인정여부")) and (v.get("인정여부") != judgements[i].인정여부
                                              or v.get("인정율") != judgements[i].인정율)
        mark = "red" if differs else None
        s.cell(f"Q{n}", v.get("인정여부") or None, fill=mark, bold=bool(v.get("인정여부")))
        s.cell(f"R{n}", v.get("인정율"), fill=mark, fmt="0%")
        s.cell(f"S{n}", f'=IF(R{n}="","",F{n}*R{n})', fmt="#,##0.0")
        s.cell(f"T{n}", v.get("사유") or None, align=LEFT, size=9)
    for k, crit_word in enumerate(("유사", "관련")):
        s.cell(f"Q{sub + k}", crit_word, fill="gray")
        s.cell(f"R{sub + k}", None, fill="gray")
        s.cell(f"S{sub + k}", f'=SUMIF({rng("Q")},"{crit_word}",{rng("S")})', fill="green", bold=True,
               fmt="#,##0.0")
        s.cell(f"T{sub + k}", None, fill="gray")
    s.box(f"Q{sub + 2}:R{sub + 2}", "총   계", fill="gray", bold=True)
    s.cell(f"S{sub + 2}", f"=S{sub}+S{sub + 1}", fill="green", bold=True, fmt="#,##0.0")
    s.cell(f"T{sub + 2}", None, fill="gray")
    if rows:
        dv1 = DataValidation(type="list", formula1='"유사,관련,불인정"', allow_blank=True)
        dv1.prompt, dv1.showInputMessage = "유사 / 관련 / 불인정", True
        dv2 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
        dv2.prompt, dv2.showInputMessage = "100%, 60%, 50%, 30% 등 (0~1)", True
        ws.add_data_validation(dv1)
        ws.add_data_validation(dv2)
        dv1.add(rng("Q"))
        dv2.add(rng("R"))

    # ── 하단 메모 ──
    notes = build_notes(meta, judgements, check, crit)
    n = sub + 4
    for k, t in enumerate(notes):
        ws.merge_cells(f"A{n + k}:O{n + k}")
        c = ws[f"A{n + k}"]
        c.value = t
        c.alignment = LEFT
        head = t.startswith("[")
        c.font = Font(name=FONT, size=9, bold=head, color="C00000" if head else "000000")
        ws.row_dimensions[n + k].height = 18

    ws.freeze_panes = "A10"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "8:9"

    out_path = Path(out_path)
    try:
        wb.save(out_path)
    except PermissionError:   # 같은 이름 파일이 엑셀에 열려 있음
        k = 2
        while (alt := out_path.with_name(f"{out_path.stem}_{k}{out_path.suffix}")).exists():
            k += 1
        wb.save(alt)
        out_path = alt
    return out_path
